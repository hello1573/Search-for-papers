import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import local_sync


class LocalSyncTests(unittest.TestCase):
    def test_newer_report_metrics_are_not_overwritten_by_older_registry(self):
        older = {'checked_at': '2026-09-26', 'if_value': 4.3, 'if_year': 2025}
        newer = {'checked_at': '2027-09-26', 'if_value': 4.5, 'if_year': 2026}
        with patch.object(local_sync.journal_metrics, 'lookup', return_value=older):
            result = local_sync.metric_columns('Example', 'journal', newer)
        self.assertEqual(result['if_value'], 4.5)
        self.assertEqual(result['if_year'], 2026)

    def test_safe_english_title_and_invalid_path_chars(self):
        self.assertEqual(local_sync.safe_filename('DEM: Sparse / Reconstruction?', 'doi'), 'DEM_ Sparse _ Reconstruction_.pdf')

    def test_docx_uses_real_hyperlinks_not_markdown_syntax(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / 'brief.docx'
            local_sync.build_docx('# Brief\n## 1. [DEM study](<https://doi.org/10.1234/example>)\n**Conclusion:** abstract only.', target, Path(directory) / 'no-template.docx')
            from docx import Document
            doc = Document(target)
            self.assertEqual(doc.paragraphs[1].text, '1. DEM study')
            self.assertIn('https://doi.org/10.1234/example', [rel.target_ref for rel in doc.part.rels.values()])
            self.assertNotIn('[DEM study]', '\n'.join(p.text for p in doc.paragraphs))

    def test_numbered_template_heading_does_not_duplicate_manual_number(self):
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        with tempfile.TemporaryDirectory() as directory:
            template = Path(directory) / 'template.docx'
            target = Path(directory) / 'brief.docx'
            source = Document()
            ppr = source.styles['Heading 1'].element.get_or_add_pPr()
            numpr = OxmlElement('w:numPr')
            numid = OxmlElement('w:numId')
            # Default Word list number 5 is a decimal list.
            numid.set(qn('w:val'), '5')
            numpr.append(numid)
            ppr.append(numpr)
            source.save(template)
            local_sync.build_docx('# Brief\n## 1. DEM study\n**本地路径：** `E:\\论文\\test.pdf`', target, template)
            doc = Document(target)
            self.assertTrue(local_sync.heading_has_number(doc, 'Heading 1'))
            self.assertEqual(doc.paragraphs[1].text, 'DEM study')
            self.assertEqual(doc.paragraphs[2].alignment, WD_ALIGN_PARAGRAPH.LEFT)
            self.assertEqual(doc.paragraphs[2].paragraph_format.first_line_indent, 0)

    def test_report_rows_keep_release_and_collection_dates_distinct(self):
        reports = [({'date': '2026-09-26', 'papers': [{
            'title': 'Sparse DEM reconstruction', 'publication_date': '2022-01-02',
            'doi': 'https://doi.org/10.1/test', 'journal': 'Journal', 'categories': ['03'],
        }]}, {})]
        rows = local_sync.rows_from_reports(reports)
        self.assertEqual(rows[0]['date'], '2022-01-02')
        self.assertEqual(rows[0]['collected'], '2026-09-26')

    def test_word_updates_keep_backups_internal_and_preserve_manual_edits(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / 'library'
            internal = Path(directory) / 'private'
            root.mkdir()
            internal.mkdir()
            output = root / 'brief.docx'
            cache = {}
            with patch.object(local_sync, 'build_docx', side_effect=lambda text, target: target.write_bytes(text.encode('utf-8'))) as builder:
                local_sync.update_docx('First version', output, '2026-09-26', cache, internal)
                local_sync.update_docx('First version', output, '2026-09-26', cache, internal)
                self.assertEqual(builder.call_count, 1)
                local_sync.update_docx('Second version', output, '2026-09-26', cache, internal)
                self.assertEqual(builder.call_count, 2)
                self.assertEqual(len(list((internal / 'backups').glob('*.docx'))), 1)
                output.write_bytes(b'User edits')
                local_sync.update_docx('Third version', output, '2026-09-26', cache, internal)
                self.assertEqual(builder.call_count, 2)
                self.assertEqual(output.read_bytes(), b'User edits')
            self.assertEqual([p.name for p in root.iterdir()], ['brief.docx'])

    def test_existing_word_is_adopted_without_rewriting_it(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'brief.docx'
            output.write_bytes(b'Existing report with user annotations')
            cache = {}
            with patch.object(local_sync, 'build_docx') as builder:
                local_sync.update_docx('Archived source', output, '2026-09-26', cache, Path(directory))
            builder.assert_not_called()
            self.assertIn('2026-09-26', cache)

    def test_internal_ledger_is_seeded_once_without_changing_legacy_file(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / 'library'
            internal = Path(directory) / 'private'
            root.mkdir()
            internal.mkdir()
            old = root / '文献收录与阅读台账.json'
            old.write_text('{"records": []}', encoding='utf-8')
            local_sync.seed_internal_records(root, internal)
            self.assertEqual((internal / old.name).read_bytes(), old.read_bytes())
            (internal / old.name).write_text('{"records": [], "marker": "current"}', encoding='utf-8')
            local_sync.seed_internal_records(root, internal)
            self.assertEqual(json.loads((internal / old.name).read_text('utf-8'))['marker'], 'current')
            self.assertNotIn('marker', json.loads(old.read_text('utf-8')))

    def test_internal_state_inside_output_directory_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with self.assertRaisesRegex(ValueError, 'outside'):
                local_sync.run(root, root, root / 'node', root / 'builder', root / 'private')

    @unittest.skipUnless(os.environ.get('TEST_NODE'), 'Requires artifact-tool Node runtime')
    def test_archived_digest_creates_matching_word_excel_and_open_pdf(self):
        report = {'schema_version': 1, 'date': '2026-09-26', 'papers': [{
            'title': 'Sparse DEM terrain reconstruction', 'publication_date': '2025-05-01',
            'doi': 'https://doi.org/10.1234/demo', 'journal': 'Example Journal',
            'relevance': '高', 'categories': ['03_缺失地形恢复_空洞填补'],
            'open_access_pdf_urls': ['https://example.org/full.pdf'],
        }]}
        markdown = '# Daily digest\n## 1. [Sparse DEM terrain reconstruction](<https://doi.org/10.1234/demo>)\n**Method:** abstract only.\n'

        class Response(io.BytesIO):
            def geturl(self):
                return 'https://example.org/full.pdf'

        def source(repo, *args):
            return json.dumps(report) if args[-1].endswith('.json') else markdown

        with tempfile.TemporaryDirectory() as directory:
            repo = Path(directory) / 'repo'
            root = Path(directory) / 'library'
            repo.mkdir()
            with patch.object(local_sync, 'report_paths', return_value=['reports/daily/2026-09-26.json']), \
                 patch.object(local_sync, 'git', side_effect=source), \
                 patch.object(local_sync, 'urlopen', side_effect=lambda *a, **kw: Response(b'%PDF-1.4\n' + b'x' * 2048 + b'\n%%EOF')):
                for _ in range(2):
                    local_sync.run(repo, root, Path(os.environ['TEST_NODE']), Path(local_sync.__file__).with_name('build_catalog.mjs'))
            self.assertTrue((root / '2026-09-26_稀疏地形重建文献简报.docx').exists())
            self.assertTrue((root / '稀疏地形重建文献总表.xlsx').exists())
            self.assertTrue((root / '论文' / '2026-09-26' / 'Sparse DEM terrain reconstruction.pdf').exists())
            self.assertEqual({p.name for p in root.iterdir()}, {'2026-09-26_稀疏地形重建文献简报.docx', '稀疏地形重建文献总表.xlsx', '论文'})
            self.assertTrue((repo / 'local-state' / '.paper_sync_state.json').exists())
            self.assertTrue((repo / 'local-state' / '.paper_catalog_index.json').exists())
            report['papers'][0]['title_zh'] = '稀疏地形重建'
            local_sync.write_json(repo / 'local-state' / '.paper_sync_state.json', {})
            rows = local_sync.rows_from_reports([(report, {})])
            source_path = repo / 'rows.json'
            source_path.write_text(json.dumps(rows), encoding='utf-8')
            local_sync.subprocess.run([os.environ['TEST_NODE'], str(Path(local_sync.__file__).with_name('build_catalog.mjs')), str(source_path), str(root / '稀疏地形重建文献总表.xlsx'), str(repo / 'local-state')], check=True, timeout=180)
            self.assertTrue(list((repo / 'local-state' / 'backups').glob('*.xlsx')))
            self.assertEqual({p.name for p in root.iterdir()}, {'2026-09-26_稀疏地形重建文献简报.docx', '稀疏地形重建文献总表.xlsx', '论文'})


if __name__ == '__main__':
    unittest.main()
