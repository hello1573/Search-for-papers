"""Mirror archived GitHub digests into the user's local research library."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


DEFAULT_ROOT = Path(r"E:\论文\知识点概述\日常文献查询记录")
DAY_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
MAX_PDF_BYTES = 40 * 1024 * 1024


def git(repo: Path, *args: str) -> str:
    result = subprocess.run(
        ["git", "-C", str(repo), *args],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=90,
        check=True,
    )
    return result.stdout


def report_paths(repo: Path) -> list[str]:
    git(repo, "fetch", "origin", "main")
    paths = git(repo, "ls-tree", "-r", "--name-only", "origin/main", "reports/daily")
    return sorted(path for path in paths.splitlines() if re.fullmatch(r"reports/daily/\d{4}-\d{2}-\d{2}\.json", path))


def safe_filename(title: str, identifier: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", title)
    name = re.sub(r"\s+", " ", name).strip(" .")[:145].rstrip(" .")
    return f"{name or hashlib.sha256(identifier.encode('utf-8')).hexdigest()[:12]}.pdf"


def fetch_open_pdf(paper: dict, folder: Path) -> Path | None:
    if paper.get("relevance") != "高":
        return None
    name = safe_filename(paper.get("title") or "", paper.get("doi") or paper.get("id") or "")
    target = folder / name
    if target.is_file() and target.stat().st_size > 1024:
        return target
    for url in paper.get("open_access_pdf_urls") or []:
        if urlparse(url).scheme != "https":
            continue
        temp_path = None
        try:
            request = Request(url, headers={"User-Agent": "sparse-terrain-digest/1.0"})
            with urlopen(request, timeout=45) as response:
                if urlparse(response.geturl()).scheme != "https":
                    continue
                with tempfile.NamedTemporaryFile(dir=folder, prefix=".pdf_", suffix=".tmp", delete=False) as output:
                    temp_path = Path(output.name)
                    count = 0
                    while chunk := response.read(128 * 1024):
                        count += len(chunk)
                        if count > MAX_PDF_BYTES:
                            raise ValueError("PDF exceeds 40 MiB limit")
                        output.write(chunk)
            with temp_path.open("rb") as source:
                header = source.read(5)
                source.seek(max(0, temp_path.stat().st_size - 2048))
                trailer = source.read()
                if header != b"%PDF-" or b"%%EOF" not in trailer or temp_path.stat().st_size < 1024:
                    raise ValueError("Response is not a complete PDF")
            os.replace(temp_path, target)
            return target
        except (OSError, ValueError) as exc:
            print(f"PDF unavailable for {paper.get('title')}: {type(exc).__name__}", file=sys.stderr)
        finally:
            if temp_path and temp_path.exists():
                temp_path.unlink()
    return None


def plain_markdown(text: str) -> str:
    text = re.sub(r"\*\*|(?<!\\)`", "", text)
    return re.sub(r"\\([\\`*_\[\]<>])", r"\1", text)


INLINE = re.compile(r"(\*\*[^*]+\*\*|\[[^]]+\]\(<https://[^>]+>\)|\[[^]]+\]\(https://[^)]+\))")
LINK = re.compile(r"^\[([^]]+)\]\(<?(https://[^)>]+)>?\)$")


def add_rich(paragraph, text: str) -> None:
    last = 0
    for match in INLINE.finditer(text):
        if match.start() > last:
            paragraph.add_run(plain_markdown(text[last:match.start()]))
        token = match.group()
        link = LINK.fullmatch(token)
        if link:
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), paragraph.part.relate_to(link.group(2), RT.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "175A77")
            properties.append(color)
            run.append(properties)
            value = OxmlElement("w:t")
            value.text = plain_markdown(link.group(1))
            run.append(value)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
        else:
            paragraph.add_run(plain_markdown(token[2:-2])).bold = True
        last = match.end()
    if last < len(text):
        paragraph.add_run(plain_markdown(text[last:]))


def build_docx(markdown: str, output: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.1)
    section.left_margin = section.right_margin = Cm(2.4)
    styles = doc.styles
    for name, size in (("Normal", 10.5), ("Title", 17), ("Heading 1", 13), ("Heading 2", 11)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(4 if name == "Normal" else 7)
    styles["Normal"].paragraph_format.line_spacing = 1.22
    for name in ("Title", "Heading 1", "Heading 2"):
        styles[name].paragraph_format.keep_with_next = True
        ppr = styles[name].element.pPr
        if ppr is not None:
            border = ppr.find(qn("w:pBdr"))
            if border is not None:
                ppr.remove(border)

    for raw in markdown.splitlines():
        line = raw.strip()
        if not line or line.startswith("@hello1573"):
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            add_rich(paragraph, line[2:])
        elif line.startswith("## "):
            add_rich(doc.add_paragraph(style="Heading 1"), line[3:])
        elif line.startswith("### "):
            add_rich(doc.add_paragraph(style="Heading 2"), line[4:])
        else:
            add_rich(doc.add_paragraph(), line)
    output.parent.mkdir(parents=True, exist_ok=True)
    temp = output.with_name("." + output.name)
    doc.save(temp)
    os.replace(temp, output)


def normalized_key(paper: dict) -> str:
    doi = (paper.get("doi") or "").lower().removeprefix("https://doi.org/")
    title = re.sub(r"\W+", "", (paper.get("title") or "").casefold())
    return doi or title


def rows_from_ledger(root: Path) -> list[dict]:
    path = root / "文献收录与阅读台账.json"
    if not path.exists():
        return []
    data = json.loads(path.read_text(encoding="utf-8"))
    rows = []
    for record in data.get("records", []):
        grades = record.get("grades") or {}
        files = record.get("local_files") or []
        rows.append({
            "date": record.get("publication_date") or record.get("publication_year") or "",
            "title": record.get("title"),
            "journal": record.get("journal"),
            "jcr": grades.get("jcr") or grades.get("status") or "未核验",
            "cas": grades.get("cas") or grades.get("status") or "未核验",
            "direction": record.get("main_direction") or "",
            "categories": ", ".join(record.get("categories") or []),
            "doi": record.get("doi"),
            "relevance": record.get("relevance") or "",
            "pdf": files[0].get("path") if files else "",
            "url": next((s for s in record.get("sources") or [] if str(s).startswith("https://")), ""),
            "collected": (record.get("pushed_dates") or [record.get("collected_at", "")[:10]])[0],
        })
    return rows


def rows_from_reports(reports: list[tuple[dict, dict[str, str]]]) -> list[dict]:
    rows = []
    for report, downloads in reports:
        for paper in report["papers"]:
            rows.append({
                "date": paper.get("publication_date"), "title": paper.get("title"),
                "journal": paper.get("journal"), "jcr": paper.get("jcr") or "未核验",
                "cas": paper.get("cas") or "未核验", "direction": paper.get("direction"),
                "categories": ", ".join(paper.get("categories") or []),
                "doi": paper.get("doi"), "relevance": paper.get("relevance"),
                "pdf": downloads.get(normalized_key(paper), ""), "url": paper.get("url"),
                "collected": report["date"],
            })
    return rows


def run(repo: Path, root: Path, node: Path, builder: Path) -> None:
    root.mkdir(parents=True, exist_ok=True)
    state_path = root / ".paper_sync_state.json"
    attempts = json.loads(state_path.read_text(encoding="utf-8")) if state_path.exists() else {}
    today = datetime.now().date().isoformat()
    reports = []
    for path in report_paths(repo):
        report = json.loads(git(repo, "show", f"origin/main:{path}"))
        day = report.get("date", "")
        if not DAY_RE.fullmatch(day) or report.get("schema_version") != 1:
            raise ValueError(f"Unrecognized report manifest: {path}")
        if path != f"reports/daily/{day}.json":
            raise ValueError(f"Report date does not match filename: {path}")
        markdown = git(repo, "show", f"origin/main:reports/daily/{day}.md")
        markdown_path = root / f"{day}_稀疏地形重建文献简报.md"
        docx_path = root / f"{day}_稀疏地形重建文献简报.docx"
        if not markdown_path.exists():
            markdown_path.write_text(markdown, encoding="utf-8")
        if not docx_path.exists():
            build_docx(markdown, docx_path)
        pdf_folder = root / "论文" / day
        pdf_folder.mkdir(parents=True, exist_ok=True)
        downloads = {}
        for paper in report["papers"]:
            key = normalized_key(paper)
            expected = pdf_folder / safe_filename(paper.get("title") or "", paper.get("doi") or paper.get("id") or "")
            if attempts.get(day + ":" + key) == today and not expected.exists():
                continue
            result = fetch_open_pdf(paper, pdf_folder)
            if result:
                downloads[key] = str(result)
            elif paper.get("relevance") == "高" and paper.get("open_access_pdf_urls"):
                attempts[day + ":" + key] = today
        reports.append((report, downloads))

    rows = rows_from_ledger(root) + rows_from_reports(reports)
    # The workbook builder appends missing records, preserving existing user edits.
    with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", suffix=".json", delete=False) as source:
        json.dump(rows, source, ensure_ascii=False)
        input_path = Path(source.name)
    try:
        workbook_path = root / "稀疏地形重建文献总表.xlsx"
        subprocess.run([str(node), str(builder), str(input_path), str(workbook_path)], check=True, timeout=180)
    finally:
        input_path.unlink(missing_ok=True)
    temp_state = state_path.with_suffix(".tmp")
    temp_state.write_text(json.dumps(attempts, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temp_state, state_path)
    print(f"已同步 {len(reports)} 期日报，文献目录 {len(rows)} 条候选记录。")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT)
    parser.add_argument("--node", type=Path, required=True)
    parser.add_argument("--builder", type=Path, default=Path(__file__).with_name("build_catalog.mjs"))
    args = parser.parse_args()
    run(args.repo, args.root, args.node, args.builder)


if __name__ == "__main__":
    main()
